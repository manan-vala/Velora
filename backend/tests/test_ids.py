import asyncio
from io import BytesIO

import openpyxl
import pytest
from pydantic import ValidationError

import geometry_processor
from algo import alns
from algo.feasibilityfinal import get_feasibility_score
from algo.lns_utils import DistanceMatrix
from algo.solver import SOLVERS
from fixtures import edge_list_from_payload, payload_from_workbook
from models import OptimizationRequest


def _payload(workbook_bytes):
    return payload_from_workbook(workbook_bytes("TestCase_TC03.xlsx"))


def _underscored(file_bytes: bytes):
    """Prefix every employee/vehicle ID in the workbook with a value containing underscores."""
    wb = openpyxl.load_workbook(BytesIO(file_bytes))
    for sheet, column, prefix in (("employees", "employee_id", "EMP_X_"), ("vehicles", "vehicle_id", "VEH_Y_")):
        ws = wb[sheet]
        col = [c.value for c in ws[1]].index(column) + 1
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=col)
            if cell.value is not None:
                cell.value = f"{prefix}{cell.value}"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


class TestValidation:
    def test_numeric_ids_are_accepted_as_strings(self, workbook_bytes):
        payload = _payload(workbook_bytes)
        payload["employees"][0]["employee_id"] = 101
        payload["vehicles"][0]["vehicle_id"] = 7
        request = OptimizationRequest(**payload)
        assert request.employees[0].employee_id == "101"
        assert request.vehicles[0].vehicle_id == "7"

    @pytest.mark.parametrize("bad_id", ["office", "Office", "", "   "])
    def test_reserved_or_empty_ids_are_rejected(self, workbook_bytes, bad_id):
        payload = _payload(workbook_bytes)
        payload["employees"][0]["employee_id"] = bad_id
        with pytest.raises(ValidationError):
            OptimizationRequest(**payload)

    def test_duplicate_ids_across_employees_and_vehicles_are_rejected(self, workbook_bytes):
        payload = _payload(workbook_bytes)
        payload["vehicles"][0]["vehicle_id"] = payload["employees"][0]["employee_id"]
        with pytest.raises(ValidationError, match="duplicated"):
            OptimizationRequest(**payload)

    def test_employees_must_share_one_office(self, workbook_bytes):
        payload = _payload(workbook_bytes)
        payload["employees"][-1]["drop_lat"] += 0.01
        with pytest.raises(ValidationError, match="drop-off"):
            OptimizationRequest(**payload)

    def test_tiny_drop_differences_are_tolerated(self, workbook_bytes):
        payload = _payload(workbook_bytes)
        payload["employees"][-1]["drop_lat"] += 0.00001
        OptimizationRequest(**payload)

    @pytest.mark.parametrize("key", ["employees", "vehicles"])
    def test_empty_lists_are_rejected(self, workbook_bytes, key):
        payload = _payload(workbook_bytes)
        payload[key] = []
        with pytest.raises(ValidationError):
            OptimizationRequest(**payload)


def test_edges_carry_their_endpoints(workbook_bytes):
    edges = edge_list_from_payload(_payload(workbook_bytes))
    for edge in edges:
        assert edge["id"] == f"{edge['from']}_{edge['to']}"


def test_distance_matrix_uses_exact_underscore_keys():
    edges = [
        {"id": "EMP_1_office", "from": "EMP_1", "to": "office", "distance_meters": 1234.0, "duration_seconds": 600.0},
        {"id": "EMP_1_2_office", "from": "EMP_1_2", "to": "office", "distance_meters": 9999.0, "duration_seconds": 60.0},
    ]
    matrix = DistanceMatrix(edges)
    assert matrix.get_dist_dur("EMP_1", "office") == (1.234, 10.0)
    assert matrix.get_dist_dur("office", "EMP_1") == (1.234, 10.0)  # reverse fallback
    assert matrix.get_dist_dur("EMP_1_2", "office") == (9.999, 1.0)


def test_alns_matrix_keeps_direction():
    alns.build_matrix([
        {"id": "a_b", "from": "a", "to": "b", "distance_meters": 1000.0, "duration_seconds": 86400.0},
        {"id": "b_a", "from": "b", "to": "a", "distance_meters": 2000.0, "duration_seconds": 43200.0},
        {"id": "a_c", "from": "a", "to": "c", "distance_meters": 3000.0, "duration_seconds": 8640.0},
    ])
    assert alns._MATRIX_DATA[("a", "b")] == (1.0, 1.0)
    assert alns._MATRIX_DATA[("b", "a")] == (2.0, 0.5)
    assert alns._MATRIX_DATA[("c", "a")] == (3.0, 0.1)  # missing reverse falls back


@pytest.mark.parametrize("solver", ["LNS", "ALNS", "VROOM"])
def test_underscore_ids_flow_through_every_solver(workbook_bytes, solver):
    file_bytes = _underscored(workbook_bytes("TestCase_TC03.xlsx"))
    payload = payload_from_workbook(file_bytes)
    assert payload["employees"][0]["employee_id"].startswith("EMP_X_")
    edges = edge_list_from_payload(payload)

    result = SOLVERS[solver](payload, edges, file_bytes, 5)
    score = get_feasibility_score(file_bytes, edges, result)

    assert score["served_count"] > 0
    known = {e["employee_id"] for e in payload["employees"]} | {v["vehicle_id"] for v in payload["vehicles"]} | {"office"}
    for vehicle in result["vehicles"]:
        assert {s["location"] for s in vehicle["route_sequence"]} <= known


def test_geometry_resolves_underscore_segments(monkeypatch, workbook_bytes):
    file_bytes = _underscored(workbook_bytes("TestCase_TC03.xlsx"))
    payload = payload_from_workbook(file_bytes)
    request = OptimizationRequest(**payload)
    emp, veh = request.employees[0], request.vehicles[0]
    requested = []

    class FakeRouteService:
        def __init__(self, max_concurrency=50):
            pass

        async def fetch_geometry_safe(self, tag, src, dst):
            requested.append((tag, src, dst))
            return tag, [[src[1], src[0]], [dst[1], dst[0]]]

        async def close(self):
            pass

    monkeypatch.setattr(geometry_processor, "RouteService", FakeRouteService)
    tags = [f"{veh.vehicle_id}_{emp.employee_id}", f"{emp.employee_id}_office"]
    schedule = {"vehicles": [{"vehicle_id": veh.vehicle_id, "routes": list(tags)}]}

    out = asyncio.run(geometry_processor.enrich_with_geometries(schedule, request))

    assert {t for t, _, _ in requested} == set(tags)
    segments = out["vehicles"][0]["route_geometry"]
    assert [s["segment_id"] for s in segments] == tags
    assert all(s["geometry"] for s in segments)
